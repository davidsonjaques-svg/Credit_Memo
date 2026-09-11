"""
Exports: a branded Excel workbook (working papers) and a Markdown DD file
that can be run through the existing md_to_pdf.py ScaleForce template.
"""
from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import DISCLAIMER, GOLD, NAVY

HEADER_FILL = PatternFill("solid", fgColor=NAVY.lstrip("#"))
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(color=NAVY.lstrip("#"), bold=True, size=14)
GOLD_FONT = Font(color=GOLD.lstrip("#"), italic=True, size=9)


def _style_sheet(ws, df: pd.DataFrame, title: str):
    ws.insert_rows(1, 3)
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = DISCLAIMER
    ws["A2"].font = GOLD_FONT
    for cell in ws[4]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, col in enumerate(df.columns, start=1):
        sample = [str(col)] + [str(v) for v in df[col].head(50).tolist()]
        width = min(max(len(s) for s in sample) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = max(width, 10)
        if pd.api.types.is_numeric_dtype(df[col]):
            fmt = "0.0%" if "pct" in str(col) or "share" in str(col) else "#,##0;(#,##0)"
            for row in ws.iter_rows(min_row=5, min_col=i, max_col=i):
                row[0].number_format = fmt
    ws.freeze_panes = "B5"


def build_excel(sheets: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            if df is None or (isinstance(df, pd.DataFrame) and df.empty):
                df = pd.DataFrame({"note": ["No items"]})
            safe = name[:31]
            df.to_excel(xw, sheet_name=safe, index=False)
            _style_sheet(xw.sheets[safe], df, name)
    return buf.getvalue()


def _fmt(v, pct=False):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "–"
    if pct:
        return f"{v:+.1%}"
    try:
        return f"{v:,.0f}"
    except (TypeError, ValueError):
        return str(v)


def _md_table(df: pd.DataFrame, cols: list, pct_cols=()) -> str:
    if df is None or df.empty:
        return "_No items._\n"
    cols = [c for c in cols if c in df.columns]
    out = ["| " + " | ".join(c.replace("_", " ").capitalize() for c in cols) + " |",
           "|" + "---|" * len(cols)]
    for _, r in df.iterrows():
        cells = []
        for c in cols:
            v = r[c]
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cells.append(_fmt(v, pct=c in pct_cols))
            else:
                cells.append("" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).replace("|", "/"))
        out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out) + "\n"


def build_markdown(extraction: dict, summary: dict | None, result: dict, review: dict | None) -> str:
    ent = extraction.get("entity") or {}
    lines = [
        f"# Financial Statement Due Diligence — {ent.get('name') or 'Entity'}",
        f"_Prepared {date.today():%d %B %Y} · Basis: {ent.get('statement_basis', 'unknown')} · "
        f"Units: {ent.get('units', 'unknown')} {ent.get('currency') or ''}_",
        "",
        f"> {DISCLAIMER}",
        "",
    ]
    if summary:
        lines += ["## 1. Executive summary", "", summary.get("business_overview", ""), ""]
        if summary.get("operating_model_indicators"):
            lines += ["**Operating model indicators**", ""] + [f"- {x}" for x in summary["operating_model_indicators"]] + [""]
        if summary.get("financial_profile_snapshot"):
            lines += ["**Financial profile**", ""] + [f"- {x}" for x in summary["financial_profile_snapshot"]] + [""]
        fa = summary.get("dd_focus_areas") or []
        if fa:
            lines += ["**Due diligence focus areas**", "",
                      "| Priority | Area | Why it matters | Line items to test |", "|---|---|---|---|"]
            for f in fa:
                lines.append(f"| {f.get('priority','')} | {f.get('area','')} | "
                             f"{f.get('why_it_matters_for_this_business','')} | "
                             f"{', '.join(f.get('line_items_to_test') or [])} |")
            lines.append("")

    reviews = {r["flag_id"]: r for r in (review or {}).get("flag_reviews", []) if r.get("flag_id")}

    def with_review(df):
        if df is None or df.empty or "flag_id" not in df:
            return df
        df = df.copy()
        df["evidence_status"] = df["flag_id"].map(lambda f: reviews.get(f, {}).get("evidence_status"))
        df["dd_question"] = df["flag_id"].map(lambda f: reviews.get(f, {}).get("dd_question"))
        return df

    mov = result["movements"]
    mov_f = mov[mov["flagged"]] if not mov.empty else mov
    lines += ["## 2. Movements above threshold", "",
              _md_table(with_review(mov_f), ["flag_id", "statement", "comparison", "line_item", "prior",
                                             "current", "pct_change", "nature", "evidence_status", "dd_question"],
                        pct_cols=("pct_change",))]
    lines += ["## 3. Unexplained line items", "",
              _md_table(with_review(result["unexplained"]),
                        ["flag_id", "statement", "line_item", "note_ref", "amount", "reasons",
                         "evidence_status", "dd_question"])]
    chk = result["integrity"]
    lines += ["## 4. Integrity checks", "",
              _md_table(with_review(chk), ["flag_id", "period", "check", "expected", "reported",
                                           "difference", "status"])]
    wc = result["wc"]
    lines += ["## 5. Receivables and payables", "",
              _md_table(wc.table, ["period", "trade_receivables", "trade_payables", "debtor_days",
                                   "creditor_days", "inventory_days", "cash_conversion_cycle"]),
              "", "**Observations**", ""]
    for o in wc.observations:
        r = reviews.get(o["flag_id"], {})
        lines.append(f"- **{o['flag_id']} · {o['area']} ({o['period']})** — {o['observation']} "
                     f"_DD query:_ {r.get('dd_question') or o['dd_query']}")
    lines += [""] + [f"_{b}_" for b in wc.basis_notes] + [""]

    if review:
        extra = review.get("additional_dd_questions") or []
        reqs = review.get("information_requests") or []
        if extra or reqs:
            lines += ["## 6. Due diligence questions and information requests", ""]
            lines += [f"- **{q.get('area','')}:** {q.get('question','')}" for q in extra]
            lines += [""] + [f"- [ ] {r}" for r in reqs] + [""]
    if extraction.get("extraction_notes"):
        lines += ["## Extraction notes", ""] + [f"- {n}" for n in extraction["extraction_notes"]] + [""]
    return "\n".join(lines)
